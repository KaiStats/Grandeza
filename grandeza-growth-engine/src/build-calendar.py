import subprocess, sys
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Pull the 30-day data out of data/content-bank.js via node so there is one source of truth.
# Path is resolved relative to this file, so `python3 src/build-calendar.py` works from repo root.
import json, os
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(REPO, "data", "content-bank.js")
raw = subprocess.check_output(
    ["node", "-e",
     "const {CONTROLS,DAYS}=require(process.argv[1]);"
     "console.log(JSON.stringify(DAYS.map(d=>({day:d.day,control:d.control,format:d.format,hook:d.hook,"
     "wt:(CONTROLS[d.control]?CONTROLS[d.control].wt:''),"
     "name:(CONTROLS[d.control]?CONTROLS[d.control].name:({MULTI:'Composite',BANDS:'Score bands',BRAND:'Founder/brand',SCAN:'The scan',SCRIPT:''}[d.control]||d.control))}))))",
     DATA])
DAYS = json.loads(raw)
OUT = os.path.join(REPO, "outputs")
os.makedirs(OUT, exist_ok=True)

NAVY="1F2A44"; GOLD="B8862B"; LIGHT="F1EFE9"; WHITE="FFFFFF"
navy_fill=PatternFill("solid",fgColor=NAVY)
gold_fill=PatternFill("solid",fgColor=GOLD)
light_fill=PatternFill("solid",fgColor=LIGHT)
thin=Side(style="thin",color="D8D2C4")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=openpyxl.Workbook()

# ---------- Calendar sheet ----------
ws=wb.active; ws.title="Calendar"
start=date(2026,7,20)  # Monday after Sat Jul 18, 2026
headers=["Day","Date","Weekday","Post time","Platform","Control #","Wt","Format","Hook / opening line","LinkedIn ✓","Facebook ✓","Status","Notes"]
widths=[6,12,11,10,20,10,6,12,60,11,11,14,26]
for c,(h,w) in enumerate(zip(headers,widths),1):
    cell=ws.cell(row=1,column=c,value=h)
    cell.font=Font(name="Arial",bold=True,color=WHITE,size=10)
    cell.fill=navy_fill; cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    cell.border=border
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=w
ws.row_dimensions[1].height=30

fmt_color={"MATH":"8A5A00","AUDIT":"1F5C4D","EXPLAINER":"1F3A5C","SCORE":"6B2A5C","STORY":"7A3520"}
for i,d in enumerate(DAYS):
    r=i+2
    dt=start+timedelta(days=i)
    row=[d["day"], dt.strftime("%b %d"), dt.strftime("%a"), "5:00 AM",
         "LinkedIn + Facebook", d["control"], d["wt"], d["format"],
         d["hook"], "", "", "Not started", ""]
    for c,val in enumerate(row,1):
        cell=ws.cell(row=r,column=c,value=val)
        cell.font=Font(name="Arial",size=9,color=fmt_color.get(d["format"],"3A3F4B") if c==8 else "3A3F4B",
                       bold=(c in (1,7,8)))
        cell.alignment=Alignment(horizontal=("left" if c in (9,13) else "center"),vertical="center",wrap_text=(c in (9,13)))
        cell.border=border
        if i%2==0: cell.fill=light_fill
    ws.row_dimensions[r].height=34

# dropdowns
dv_status=DataValidation(type="list",formula1='"Not started,Drafted,Scheduled,Posted,Skipped"',allow_blank=True)
dv_check=DataValidation(type="list",formula1='"✓,—"',allow_blank=True)
ws.add_data_validation(dv_status); ws.add_data_validation(dv_check)
dv_status.add(f"L2:L{len(DAYS)+1}")
dv_check.add(f"J2:J{len(DAYS)+1}"); dv_check.add(f"K2:K{len(DAYS)+1}")
ws.freeze_panes="A2"
ws.auto_filter.ref=f"A1:M{len(DAYS)+1}"

# ---------- Legend sheet ----------
lg=wb.create_sheet("Legend")
lg.column_dimensions["A"].width=16; lg.column_dimensions["B"].width=64
def hrow(ws_,r,text):
    c=ws_.cell(row=r,column=1,value=text); c.font=Font(name="Arial",bold=True,color=WHITE,size=11); c.fill=navy_fill
    ws_.cell(row=r,column=2).fill=navy_fill
def kv(ws_,r,k,v):
    a=ws_.cell(row=r,column=1,value=k); a.font=Font(name="Arial",bold=True,size=10,color="3A3F4B"); a.alignment=Alignment(vertical="top")
    b=ws_.cell(row=r,column=2,value=v); b.font=Font(name="Arial",size=10,color="3A3F4B"); b.alignment=Alignment(wrap_text=True,vertical="top")

lg.cell(row=1,column=1,value="GRANDEZA 30-DAY CALENDAR — LEGEND").font=Font(name="Arial",bold=True,size=13,color=NAVY)
r=3
hrow(lg,r,"How to use"); r+=1
for k,v in [
 ("Fill in","Only the columns you touch: the two platform checkmarks (J, K), Status (L), and Notes (M). Everything else is set."),
 ("Post time","5:00 AM daily. Batch-schedule a week ahead (LinkedIn native scheduler; Meta Business Suite for Facebook — both free)."),
 ("Status","Drop-down: Not started → Drafted → Scheduled → Posted (or Skipped)."),
 ("Checkmarks","J/K drop-down (✓ / —) so you can see at a glance which platform a day has gone live on."),
]:
    kv(lg,r,k,v); r+=1
r+=1
hrow(lg,r,"Format codes"); r+=1
for k,v in [
 ("MATH","Text-on-screen exposure stacking — the dollar figures. Highest-converting hooks."),
 ("AUDIT","Live or screen-recorded look at a real bot/site failing a control."),
 ("EXPLAINER","Talking-head or voiceover teaching one control in plain English."),
 ("SCORE","Run the free scan live and narrate the number as it populates."),
 ("STORY","Founder / recon narrative (gas-leak call, Clark County investigator angle)."),
]:
    kv(lg,r,k,v); r+=1
r+=1
hrow(lg,r,"The 4 critical controls (auto-fails)"); r+=1
for k,v in [
 ("03","TCPA Outbound (weight 9) — $500–$1,500/text, no class-action cap."),
 ("09","Recording Consent (weight 8) — NRS 200.620, Nevada category D felony."),
 ("01","AI Identity Disclosure (weight 8) — FTC deceptive-practices exposure today."),
 ("14","Audit Trail (weight 8) — no log, no defense."),
]:
    kv(lg,r,k,v); r+=1
r+=1
hrow(lg,r,"Legal note"); r+=1
kv(lg,r,"Verified","No 'SB 149' / 'Oct 1 deadline' anywhere — that claim is not real. Spine = federal TCPA, NRS 200.620, NRS 603A, FTC §5, multi-state disclosure wave. Get a one-time NV attorney review before scaling the paid binder."); r+=1

for row in lg.iter_rows(min_row=1,max_row=r,min_col=1,max_col=2):
    for cell in row:
        cell.border=Border(bottom=Side(style="hair",color="E6E1D6"))

wb.save(os.path.join(OUT,"Grandeza_30Day_Content_Calendar.xlsx"))
print("calendar written")
